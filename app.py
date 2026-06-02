import streamlit as st
import openai
import os
import io
import re
import requests

# ── Page config ────────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="Ask BIONEXT",
    page_icon="🌿",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ── Branding / CSS ─────────────────────────────────────────────────────────────
st.markdown("""
<style>
  :root {
    --blue:  #1A6B7C;
    --green: #8A9A3A;
    --light: #F0F5F5;
  }
  .bionext-header {
    background-color: var(--blue);
    padding: 1rem 2rem;
    border-radius: 8px;
    margin-bottom: 1.5rem;
    display: flex;
    align-items: center;
    gap: 1.2rem;
  }
  .user-msg {
    background-color: var(--blue);
    color: white;
    padding: 0.8rem 1.1rem;
    border-radius: 12px 12px 2px 12px;
    margin: 0.5rem 0 0.5rem 15%;
    font-size: 0.95rem;
  }
  .assistant-msg {
    background-color: white;
    border: 1.5px solid #e0e0e0;
    padding: 0.9rem 1.1rem;
    border-radius: 12px 12px 12px 2px;
    margin: 0.5rem 15% 0.5rem 0;
    font-size: 0.95rem;
    line-height: 1.6;
  }
  [data-testid="stSidebar"] {
    background-color: var(--light);
    border-right: 3px solid var(--green);
  }
  .sidebar-title {
    color: var(--blue);
    font-weight: 700;
    font-size: 1rem;
    margin-bottom: 0.4rem;
  }
  .stTextInput input {
    border: 2px solid var(--blue) !important;
    border-radius: 8px !important;
  }
  .stButton button {
    background-color: var(--blue) !important;
    color: white !important;
    border: none !important;
    border-radius: 8px !important;
    font-weight: 600 !important;
    padding: 0.5rem 1.5rem !important;
  }
  .stButton button:hover {
    background-color: var(--green) !important;
    color: var(--blue) !important;
  }
  footer {visibility: hidden;}
</style>
""", unsafe_allow_html=True)

# ── Header ─────────────────────────────────────────────────────────────────────
logo_url = "[raw.githubusercontent.com](https://raw.githubusercontent.com/PM-Oppla/ask-bionext/main/bionext-logo.png)"
st.markdown(f"""
<div class="bionext-header">
  <img src="{logo_url}" height="55" style="margin-right:1rem; flex-shrink:0;">
  <div>
    <h1 style="color:white; margin:0; font-size:1.8rem; font-weight:700;">Ask BIONEXT</h1>
    <p style="color:#8A9A3A; margin:0; font-size:0.95rem;">Have a conversation with BIONEXT research · The Biodiversity Nexus: Transformative Change for Sustainability</p>
  </div>
</div>
""", unsafe_allow_html=True)

# ── API key ────────────────────────────────────────────────────────────────────
api_key = st.secrets.get("OPENAI_API_KEY", os.environ.get("OPENAI_API_KEY", ""))
client = openai.OpenAI(api_key=api_key)

# ── Session state ──────────────────────────────────────────────────────────────
if "messages" not in st.session_state:
    st.session_state.messages = []
if "documents" not in st.session_state:
    st.session_state.documents = {}
if "drive_loaded" not in st.session_state:
    st.session_state.drive_loaded = False

# ── Google Drive folder ID ─────────────────────────────────────────────────────
DRIVE_FOLDER_ID = "1kYY0erFzaR5UNDG_px-gO442GUlHZYZo"

# ── Fetch file list from public Google Drive folder ────────────────────────────
def get_drive_files(folder_id):
    """Scrape file IDs from a public Google Drive folder page."""
    url = f"[drive.google.com](https://drive.google.com/drive/folders/{folder_id})"
    headers = {"User-Agent": "Mozilla/5.0"}
    try:
        r = requests.get(url, headers=headers, timeout=15)
        # Extract file IDs using regex pattern from Drive page source
        ids_names = re.findall(r'"([a-zA-Z0-9_-]{33})","([^"]+\.[a-zA-Z]{2,5})"', r.text)
        # Deduplicate
        seen = set()
        files = []
        for fid, fname in ids_names:
            if fid not in seen:
                seen.add(fid)
                files.append({"id": fid, "name": fname})
        return files
    except Exception as e:
        return []

# ── Download file from Google Drive ───────────────────────────────────────────
def download_drive_file(file_id):
    """Download a file from Google Drive by ID."""
    url = f"[drive.google.com](https://drive.google.com/uc?export=download&id={file_id})"
    try:
        session = requests.Session()
        r = session.get(url, timeout=30, stream=True)
        # Handle virus scan warning for larger files
        for key, value in r.cookies.items():
            if key.startswith("download_warning"):
                r = session.get(url, params={"confirm": value}, timeout=30)
                break
        return io.BytesIO(r.content)
    except Exception:
        return None

# ── Extract text from file bytes ──────────────────────────────────────────────
def extract_text(file_bytes, filename):
    name = filename.lower()
    try:
        if name.endswith(".pdf"):
            import pypdf
            reader = pypdf.PdfReader(file_bytes)
            return "\n\n".join(p.extract_text() or "" for p in reader.pages)
        elif name.endswith(".docx"):
            import docx
            doc = docx.Document(file_bytes)
            return "\n".join(p.text for p in doc.paragraphs)
        elif name.endswith(".xlsx"):
            import openpyxl
            wb = openpyxl.load_workbook(file_bytes, data_only=True)
            lines = []
            for ws in wb.worksheets:
                lines.append(f"[Sheet: {ws.title}]")
                for row in ws.iter_rows(values_only=True):
                    lines.append("\t".join(str(c) if c is not None else "" for c in row))
            return "\n".join(lines)
        elif name.endswith((".txt", ".csv", ".md")):
            return file_bytes.read().decode("utf-8", errors="ignore")
    except Exception:
        return ""
    return ""

# ── Smart retrieval ────────────────────────────────────────────────────────────
def get_relevant_docs(question, documents, top_n=5):
    stop_words = {'what','how','why','when','where','who','which','is','are','the',
                  'a','an','in','on','at','to','for','of','and','or','do','does',
                  'did','can','could','tell','me','about','any','please'}
    question_words = set(re.sub(r'[^\w\s]', '', question.lower()).split()) - stop_words
    scores = {}
    for fname, text in documents.items():
        text_lower = text.lower()
        score = sum(text_lower.count(w) for w in question_words)
        score += sum(10 for w in question_words if w in fname.lower())
        scores[fname] = score
    sorted_docs = sorted(scores.items(), key=lambda x: x[1], reverse=True)
    return [f for f, s in sorted_docs[:top_n]]

# ── Build context ──────────────────────────────────────────────────────────────
def build_context(question):
    if not st.session_state.documents:
        return "", []
    relevant = get_relevant_docs(question, st.session_state.documents)
    parts = []
    for fname in relevant:
        text = st.session_state.documents[fname]
        snippet = text[:8000] if len(text) > 8000 else text
        parts.append(f"=== SOURCE: {fname} ===\n{snippet}\n")
    return "\n".join(parts), relevant

# ── System prompt ──────────────────────────────────────────────────────────────
SYSTEM_PROMPT = """You are 'Ask BIONEXT', an AI research assistant for the BIONEXT project
(The Biodiversity Nexus: Transformative Change for Sustainability) — a European Union research
project exploring how biodiversity interconnects with climate, food, water, energy, transport,
and health.

Answer questions by drawing EXCLUSIVELY on the source documents provided.
Do not use any outside knowledge. If the answer is not in the sources, say so clearly.

For every answer:
1. Ground every claim in the source documents
2. End with: 📄 Sources: [filename1], [filename2]
3. Be clear and accessible — users may be policymakers, researchers, or public
4. If unanswerable from sources: "I don't have information on that in the current
   BIONEXT documents. You may find more at bionext-project.eu"

Source documents:
{context}
"""

# ── Sidebar ────────────────────────────────────────────────────────────────────
with st.sidebar:
    st.markdown('<div class="sidebar-title">📂 BIONEXT Sources</div>', unsafe_allow_html=True)

    if not st.session_state.drive_loaded:
        if st.button("🔄 Load BIONEXT Documents"):
            with st.spinner("Connecting to BIONEXT document library..."):
                files = get_drive_files(DRIVE_FOLDER_ID)
                if files:
                    progress = st.progress(0)
                    loaded = 0
                    errors = 0
                    for i, f in enumerate(files):
                        name = f["name"]
                        if not any(name.lower().endswith(ext) for ext in
                                   [".pdf",".docx",".xlsx",".txt",".csv",".md"]):
                            continue
                        file_bytes = download_drive_file(f["id"])
                        if file_bytes:
                            text = extract_text(file_bytes, name)
                            if text.strip():
                                st.session_state.documents[name] = text
                                loaded += 1
                            else:
                                errors += 1
                        progress.progress(min((i + 1) / max(len(files), 1), 1.0))
                    st.session_state.drive_loaded = True
                    if loaded > 0:
                        st.success(f"✓ {loaded} documents loaded")
                    else:
                        st.error("No documents could be loaded. The folder may not be fully public.")
                else:
                    st.error("Could not read the folder. Please ensure sharing is set to 'Anyone with the link'.")
    else:
        st.success(f"✓ {len(st.session_state.documents)} documents loaded")

    if st.session_state.documents:
        st.markdown("---")
        st.markdown('<div class="sidebar-title">Loaded sources</div>', unsafe_allow_html=True)
        for name in sorted(st.session_state.documents.keys()):
            st.markdown(f"📄 `{name}`")

    st.markdown("---")
    st.markdown('<div class="sidebar-title">Or upload files manually</div>', unsafe_allow_html=True)
    uploaded = st.file_uploader(
        "Upload",
        type=["pdf","txt","docx","xlsx","csv","md"],
        accept_multiple_files=True,
        label_visibility="collapsed"
    )
    if uploaded:
        for f in uploaded:
            if f.name not in st.session_state.documents:
                text = extract_text(io.BytesIO(f.read()), f.name)
                if text:
                    st.session_state.documents[f.name] = text
                    st.success(f"✓ {f.name}")

    st.markdown("---")
    if st.button("🗑 Clear conversation"):
        st.session_state.messages = []
        st.rerun()

# ── Chat display ───────────────────────────────────────────────────────────────
if not st.session_state.messages:
    st.markdown("""
    <div class="assistant-msg">
    👋 Welcome! I'm <strong>Ask BIONEXT</strong> — your guide to BIONEXT project research.<br><br>
    Click <strong>'Load BIONEXT Documents'</strong> in the sidebar to load the document library,
    then ask me anything about the project's findings, methods, or outputs.
    </div>
    """, unsafe_allow_html=True)

for msg in st.session_state.messages:
    if msg["role"] == "user":
        st.markdown(f'<div class="user-msg">{msg["content"]}</div>', unsafe_allow_html=True)
    else:
        st.markdown(f'<div class="assistant-msg">{msg["content"]}</div>', unsafe_allow_html=True)

# ── Input ──────────────────────────────────────────────────────────────────────
st.markdown("---")
col1, col2 = st.columns([5, 1])
with col1:
    user_input = st.text_input(
        "Your question",
        placeholder="e.g. What are the key findings on biodiversity and food systems?",
        label_visibility="collapsed",
        key="input"
    )
with col2:
    send = st.button("Ask →")

# ── Response logic ─────────────────────────────────────────────────────────────
if send and user_input.strip():
    if not st.session_state.documents:
        st.warning("⚠️ Please load the BIONEXT documents first using the button in the sidebar.")
    elif not api_key:
        st.error("⚠️ No OpenAI API key found.")
    else:
        st.session_state.messages.append({"role": "user", "content": user_input})
        context, _ = build_context(user_input)
        system = SYSTEM_PROMPT.format(context=context)
        with st.spinner("Searching BIONEXT research..."):
            try:
                response = client.chat.completions.create(
                    model="gpt-4o-mini",
                    messages=[
                        {"role": "system", "content": system},
                        *[{"role": m["role"], "content": m["content"]}
                          for m in st.session_state.messages]
                    ],
                    temperature=0.2,
                    max_tokens=1200
                )
                answer = response.choices[0].message.content
                st.session_state.messages.append({"role": "assistant", "content": answer})
            except Exception as e:
                st.error(f"API error: {e}")
        st.rerun()
